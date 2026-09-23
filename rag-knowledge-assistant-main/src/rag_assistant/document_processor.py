from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"}


class DocumentProcessingError(ValueError):
    """Raised when a document cannot be parsed or validated."""


@dataclass(frozen=True)
class ParsedDocument:
    source_id: str
    content: str
    metadata: dict[str, Any]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_source_id(filename: str) -> str:
    stem = Path(filename).stem.strip() or "document"
    value = re.sub(r"[^a-zA-Z0-9._-]+", "-", stem).strip(".-")
    return value[:160] or "document"


def _extract_pdf(data: bytes) -> tuple[str, dict[int, str]]:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise DocumentProcessingError("PDF support requires the 'pypdf' package.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        pages: dict[int, str] = {}
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages[index] = text
        marked = "\n\n".join(f"[PAGE {page}]\n{text}" for page, text in pages.items())
        return marked, pages
    except Exception as exc:
        raise DocumentProcessingError("Could not extract text from the PDF document.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise DocumentProcessingError("DOCX support requires the 'python-docx' package.") from exc

    try:
        document = Document(io.BytesIO(data))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        return "\n\n".join(blocks)
    except Exception as exc:
        raise DocumentProcessingError("Could not extract text from the DOCX document.") from exc


def parse_document(filename: str, data: bytes, source_id: str | None = None) -> ParsedDocument:
    if not data:
        raise DocumentProcessingError("The uploaded document is empty.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentProcessingError(f"Unsupported file type '{suffix or 'unknown'}'. Supported: {supported}.")

    pages: dict[int, str] = {}
    if suffix == ".pdf":
        content, pages = _extract_pdf(data)
    elif suffix == ".docx":
        content = _extract_docx(data)
    else:
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError("Text documents must be UTF-8 encoded.") from exc

    content = content.replace("\x00", " ").strip()
    if not content:
        raise DocumentProcessingError("No readable text was found in the document.")

    digest = hashlib.sha256(data).hexdigest()
    document_id = source_id.strip() if source_id else f"{_safe_source_id(filename)}-{digest[:10]}"
    metadata: dict[str, Any] = {
        "filename": filename,
        "file_type": suffix.lstrip("."),
        "content_hash": digest,
        "ingested_at": _now(),
        "updated_at": _now(),
        "page_count": len(pages),
    }

    return ParsedDocument(source_id=document_id, content=content, metadata=metadata)
